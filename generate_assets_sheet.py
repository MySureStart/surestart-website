import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CDN = "https://pub-dcb79d88e2ec4567824453cc853f55f3.r2.dev"

rows = [
    # (Source File, Type, URL, Alt Text)

    # ── CSS Backgrounds ──────────────────────────────────────────────────────
    ("assets/css/about-us.css",       "CSS background", CDN + "/images/heroes/about-us-hero.jpg",                                  ""),
    ("assets/css/contact-us.css",     "CSS background", CDN + "/images/heroes/contact-us-hero.jpg",                                ""),
    ("assets/css/higher-ed.css",      "CSS background", CDN + "/images/heroes/higher-ed-hero.jpg",                                 ""),
    ("assets/css/impact-stories.css", "CSS background", CDN + "/images/heroes/impact-stories-hero.jpg",                            ""),
    ("assets/css/k12.css",            "CSS background", CDN + "/images/heroes/k-12-header.png",                                    ""),
    ("assets/css/styles.css",         "CSS background", CDN + "/images/icons/pen-underline.svg",                                   ""),
    ("assets/css/about-us.css",       "CSS background", CDN + "/images/icons/pen-underline.svg",                                   ""),
    ("assets/css/contact-us.css",     "CSS background", CDN + "/images/icons/pen-underline.svg",                                   ""),
    ("assets/css/higher-ed.css",      "CSS background", CDN + "/images/icons/pen-underline.svg",                                   ""),
    ("assets/css/k12.css",            "CSS background", CDN + "/images/icons/pen-underline.svg",                                   ""),
    ("assets/css/students.css",       "CSS background", CDN + "/images/icons/pen-underline.svg",                                   ""),
    ("assets/css/impact-stories.css", "CSS background", CDN + "/images/icons/pen-underline.svg",                                   ""),
    ("assets/css/vibe-lab.css",       "CSS background", CDN + "/images/icons/pen-underline.svg",                                   ""),
    ("assets/css/about-us.css",       "CSS background", CDN + "/images/misc/fortune-article-bg.png",                               ""),
    ("assets/css/about-us.css",       "CSS background", CDN + "/images/media/pioneers-of-ai.jpg",                                  ""),
    ("assets/css/impact-stories.css", "CSS background", CDN + "/images/misc/world-map.png",                                        ""),
    ("assets/css/impact-stories.css", "CSS background", CDN + "/images/icons/globe.png",                                           ""),
    ("assets/css/students.css",       "CSS background", CDN + "/images/success-stories/k12/AdobeStock_956183976.jpeg",             ""),

    # ── Inline CSS Backgrounds — index.html ──────────────────────────────────
    ("index.html", "Inline CSS background", CDN + "/images/programs/1.png", ""),
    ("index.html", "Inline CSS background", CDN + "/images/programs/2.png", ""),
    ("index.html", "Inline CSS background", CDN + "/images/programs/3.png", ""),
    ("index.html", "Inline CSS background", CDN + "/images/programs/4.png", ""),

    # ── Inline CSS Backgrounds — k12/index.html ──────────────────────────────
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/1.png",            ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/2.png",            ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/3.png",            ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/4.png",            ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/5.png",            ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/6.png",            ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/81.png",           ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/91.png",           ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/ai-gaming.jpg",    ""),
    ("k12/index.html", "Inline CSS background", CDN + "/images/programs/ai-creativity.jpg",""),

    # ── Inline CSS Backgrounds — for-students/index.html ─────────────────────
    ("for-students/index.html", "Inline CSS background", CDN + "/images/success-stories/k12/AdobeStock_572964027_Preview.jpeg", ""),
    ("for-students/index.html", "Inline CSS background", CDN + "/images/success-stories/k12/mit-futuremakers.jpg",             ""),
    ("for-students/index.html", "Inline CSS background", CDN + "/images/success-stories/k12/AdobeStock_575673643.jpeg",        ""),
    ("for-students/index.html", "Inline CSS background", CDN + "/images/success-stories/k12/branksome-hall.jpg",               ""),

    # ── <img> — index.html ───────────────────────────────────────────────────
    ("index.html", "<img>", CDN + "/images/logos/surestart/surestart-logo.png",      "SureStart"),
    ("index.html", "<img>", CDN + "/images/logos/universities/harvard.png",           "Harvard University"),
    ("index.html", "<img>", CDN + "/images/logos/universities/uc-berkeley.png",       "UC Berkeley"),
    ("index.html", "<img>", CDN + "/images/logos/universities/princeton.png",         "Princeton University"),
    ("index.html", "<img>", CDN + "/images/logos/universities/stanford.png",          "Stanford University"),
    ("index.html", "<img>", CDN + "/images/logos/universities/dartmouth.png",         "Dartmouth"),
    ("index.html", "<img>", CDN + "/images/logos/universities/cuny.png",              "The City University of New York"),
    ("index.html", "<img>", CDN + "/images/logos/universities/georgia-tech.png",      "Georgia Tech"),
    ("index.html", "<img>", CDN + "/images/logos/companies/microsoft.png",            "Microsoft"),
    ("index.html", "<img>", CDN + "/images/logos/companies/nasa.png",                 "NASA"),
    ("index.html", "<img>", CDN + "/images/logos/companies/meta.png",                 "Meta"),
    ("index.html", "<img>", CDN + "/images/logos/companies/consumer-reports.png",     "Consumer Reports"),
    ("index.html", "<img>", CDN + "/images/logos/companies/red-hat.png",              "Red Hat"),
    ("index.html", "<img>", CDN + "/images/logos/companies/salesforce.png",           "Salesforce"),
    ("index.html", "<img>", CDN + "/images/logos/companies/linkedin.png",             "Linkedin"),
    ("index.html", "<img>", CDN + "/images/projects/chalkit.png",                     "Chalk It Demo Video Thumbnail"),
    ("index.html", "<img>", CDN + "/images/projects/nourisync.png",                   "NouriSync Demo Video Thumbnail"),
    ("index.html", "<img>", CDN + "/images/projects/hand2hand.png",                   "Hand2Hand Demo Video Thumbnail"),
    ("index.html", "<img>", CDN + "/images/team/students/ammran-mohamed.png",         "Ammran H Mohamed"),
    ("index.html", "<img>", CDN + "/images/team/students/vitaliy-stepanov.png",       "Vitaliy Stephanov"),
    ("index.html", "<img>", CDN + "/images/team/students/alexa-urrea.png",            "Alexa Urrea"),
    ("index.html", "<img>", CDN + "/images/team/students/jonathan-williams.png",      "Jonathan Williams"),
    ("index.html", "<img>", CDN + "/images/team/students/netra-ramesh.png",           "Netra Ramesh"),
    ("index.html", "<img>", CDN + "/images/team/students/jed-rendo-m.png",            "Jed Rendo Margarcia"),
    ("index.html", "<img>", CDN + "/images/team/students/ashmita-kumar.png",          "Ashmita Kumar"),
    ("index.html", "<img>", CDN + "/images/team/students/ashna-khetan.png",           "Ashna Khetan"),

    # ── <video> — index.html ─────────────────────────────────────────────────
    ("index.html", "<video>", CDN + "/videos/heroes/home-hero-loop.mp4", ""),

    # ── <img> — about/index.html ─────────────────────────────────────────────
    ("about/index.html", "<img>", CDN + "/images/logos/surestart/surestart-logo.png",                                        "SureStart"),
    ("about/index.html", "<img>", CDN + "/images/heroes/our-story.jpg",                                                      "SureStart students working together"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/fortune.png",                                                   "Fortune"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/mit-tech-review.png",                                           "MIT Technology Review"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/nbc-learn.png",                                                 "NBC Learn"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/the-atlantic.png",                                              "The Atlantic"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/colby-news.webp",                                               "Colby News"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/thrive-global.png",                                             "Thrive Global"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/fast-company.png",                                              "Fast Company"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/women-in-tech-2021.jpeg",                                       "WomenTech Network Global AI Inclusion Award"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/venturebeat.svg",                                               "VentureBeat Woman in AI"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/women-in-ai-awards.png",                                        "Women in AI North America Award"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/nasdaq.svg",                                                    "NASDAQ 25 Leading Women in AI"),
    ("about/index.html", "<img>", CDN + "/images/logos/press/asu-gsv.webp",                                                  "ASU+GSV Innovator of Color"),
    ("about/index.html", "<img>", CDN + "/images/logos/partners/iste.png",                                                   "ISTE 20 to Watch Award"),
    ("about/index.html", "<img>", CDN + "/images/logos/partners/pioneers-ai.png",                                            "Company & Partner Logos"),
    ("about/index.html", "<img>", CDN + "/images/media/futurized.jpeg",                                                      "Futurized Podcast"),
    ("about/index.html", "<img>", CDN + "/images/media/getting-smart.webp",                                                  "Taniya Mishra and Amanda Stent on Getting Smart Podcast"),
    ("about/index.html", "<img>", CDN + "/images/media/human-centric-ai.webp",                                               "Taniya Mishra on Human Centric AI Podcast"),
    ("about/index.html", "<img>", CDN + "/images/media/technolotea.jpeg",                                                    "Technolotea Podcast"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/taniya-mishra.jpg",                                         "Dr. Taniya Mishra"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/taniya-cropped.png",                                        "Taniya Mishra"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/ashleen.jpg",                                               "Ashleen Sullivan"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/nicholas.jpg",                                              "Nicholas Glading"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/kim.jpg",                                                   "Kim Parnell"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/betty.jpg",                                                 "Betty Wong"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/ericka-corral.png",                                         "Ericka Corral"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/robin.png",                                                 "Robin Kondrup"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/isabela.jpg",                                               "Isabela Daudt"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/LaurenHeadshotHD.jpeg",                                     "Lauren Kilefner"),
    ("about/index.html", "<img>", CDN + "/images/team/leadership/WhatsApp%20Image%202026-04-17%20at%2017.18.09.jpeg",        "Islam Zidan"),
    ("about/index.html", "<img>", CDN + "/images/team/mentors/cecilia-dones.jpg",                                            "Dr. Cecilia Dones"),
    ("about/index.html", "<img>", CDN + "/images/team/mentors/leonardo-neves.jpeg",                                          "Leonardo Neves"),
    ("about/index.html", "<img>", CDN + "/images/team/mentors/safinah-ali.jpeg",                                             "Dr. Safinah Ali"),
    ("about/index.html", "<img>", CDN + "/images/team/mentors/sarah-ward.jpg",                                               "Sarah Ward"),

    # ── <img> — contact/index.html ───────────────────────────────────────────
    ("contact/index.html", "<img>", CDN + "/images/logos/surestart/surestart-logo.png", "SureStart"),

    # ── <img> — vibe-lab/index.html ──────────────────────────────────────────
    ("vibe-lab/index.html", "<img>", CDN + "/images/logos/surestart/surestart-logo.png",  "SureStart"),
    ("vibe-lab/index.html", "<img>", CDN + "/images/heroes/vibe-lab-hero.png",             "Students collaborating on AI projects in Vibe Lab"),
    ("vibe-lab/index.html", "<img>", CDN + "/images/misc/vibe-lab-details.png",            "Vibe Lab program details and curriculum overview"),

    # ── <img> — k12/index.html ───────────────────────────────────────────────
    ("k12/index.html", "<img>", CDN + "/images/logos/surestart/surestart-logo.png",          "SureStart"),
    ("k12/index.html", "<img>", CDN + "/images/icons/early-access.png",                      "Early Access Icon"),
    ("k12/index.html", "<img>", CDN + "/images/icons/future-proof.png",                      "Future-Proof Icon"),
    ("k12/index.html", "<img>", CDN + "/images/icons/personalized-mentorship.png",           "Personalized Mentorship Icon"),
    ("k12/index.html", "<img>", CDN + "/images/icons/essential-skills.png",                  "Essential Skills Icon"),
    ("k12/index.html", "<img>", CDN + "/images/icons/global-reach.png",                      "Global Reach Icon"),
    ("k12/index.html", "<img>", CDN + "/images/icons/empowering-teachers.png",               "Empowering Teachers Icon"),
    ("k12/index.html", "<img>", CDN + "/images/logos/universities/harvard.png",              "Harvard University"),
    ("k12/index.html", "<img>", CDN + "/images/logos/universities/uc-berkeley.png",          "UC Berkeley"),
    ("k12/index.html", "<img>", CDN + "/images/logos/universities/princeton.png",            "Princeton University"),
    ("k12/index.html", "<img>", CDN + "/images/logos/universities/stanford.png",             "Stanford University"),
    ("k12/index.html", "<img>", CDN + "/images/logos/universities/dartmouth.png",            "Dartmouth"),
    ("k12/index.html", "<img>", CDN + "/images/logos/universities/cuny.png",                 "The City University of New York"),
    ("k12/index.html", "<img>", CDN + "/images/logos/universities/georgia-tech.png",         "Georgia Tech"),
    ("k12/index.html", "<img>", CDN + "/images/logos/partners/mit-raise.png",                "MIT Media Lab"),
    ("k12/index.html", "<img>", CDN + "/images/logos/partners/xcl.png",                      "XCL Education"),
    ("k12/index.html", "<img>", CDN + "/images/logos/partners/branksome-hall.png",           "Branksome Hall"),
    ("k12/index.html", "<img>", CDN + "/images/logos/partners/xcl.png",                      "XCL World Academy"),
    ("k12/index.html", "<img>", CDN + "/images/icons/graduation-hat.png",                    "Graduation Hat"),
    ("k12/index.html", "<img>", CDN + "/images/success-stories/k12/xcl.jpg",                 "XCL World Academy"),
    ("k12/index.html", "<img>", CDN + "/images/success-stories/k12/mit-futuremakers.jpg",    "MIT FutureMakers"),
    ("k12/index.html", "<img>", CDN + "/images/success-stories/k12/branksome-hall.jpg",      "Branksome Hall"),

    # ── <img> — for-universities/index.html ──────────────────────────────────
    ("for-universities/index.html", "<img>", CDN + "/images/logos/surestart/surestart-logo.png",          "SureStart"),
    ("for-universities/index.html", "<img>", CDN + "/images/icons/customized-ai-programs.png",            "Customized AI Programs"),
    ("for-universities/index.html", "<img>", CDN + "/images/icons/plug-and-play-curriculum.png",          "Plug-and-Play Curriculum"),
    ("for-universities/index.html", "<img>", CDN + "/images/icons/durable-skill-integration.png",         "Durable Skill Integration"),
    ("for-universities/index.html", "<img>", CDN + "/images/icons/consulting-support.png",                "Consulting Support"),
    ("for-universities/index.html", "<img>", CDN + "/images/icons/professional-network-access.png",       "Professional Network Access"),
    ("for-universities/index.html", "<img>", CDN + "/images/icons/flexible-implementation.png",           "Flexible Implementation"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/companies/microsoft.png",               "Microsoft"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/companies/nasa.png",                    "NASA"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/companies/meta.png",                    "Meta"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/companies/consumer-reports.png",        "Consumer Reports"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/companies/red-hat.png",                 "Red Hat"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/companies/salesforce.png",              "Salesforce"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/companies/linkedin.png",                "Linkedin"),
    ("for-universities/index.html", "<img>", CDN + "/images/success-stories/higher-ed/pas4ai.png",        "PAS4AI Partnership"),
    ("for-universities/index.html", "<img>", CDN + "/images/success-stories/higher-ed/ai-trailblazers.jpg", "SureStart AI Trailblazers"),
    ("for-universities/index.html", "<img>", CDN + "/images/success-stories/higher-ed/break-through-ai.jpg", "Break Through Tech"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/partners/colby.png",                    "Colby College"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/partners/penn-state.png",               "Penn State"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/partners/mit-raise.png",                "MIT Media Lab"),
    ("for-universities/index.html", "<img>", CDN + "/images/logos/universities/cmu.png",                  "Carnegie Mellon University"),

    # ── <img> — impact-stories/index.html ────────────────────────────────────
    ("impact-stories/index.html", "<img>", CDN + "/images/logos/surestart/surestart-logo.png",      "SureStart"),
    ("impact-stories/index.html", "<img>", CDN + "/images/icons/graduation-hat.png",                "Graduation Hat"),
    ("impact-stories/index.html", "<img>", CDN + "/images/misc/lili-piesanen.png",                  "Lili Piesanen"),
    ("impact-stories/index.html", "<img>", CDN + "/images/logos/universities/mcgill.png",           "McGill University"),
    ("impact-stories/index.html", "<img>", CDN + "/images/misc/jonathan-williams.jpg",              "Jonathan Williams"),
    ("impact-stories/index.html", "<img>", CDN + "/images/logos/universities/princeton.png",        "Princeton University"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/leadership/ericka-corral.png",       "Ericka Corral"),
    ("impact-stories/index.html", "<img>", CDN + "/images/logos/universities/denver.png",           "University of Denver"),
    ("impact-stories/index.html", "<img>", CDN + "/images/misc/annika.jpeg",                        "Annika Sachdeva"),
    ("impact-stories/index.html", "<img>", CDN + "/images/logos/universities/mit.png",              "massachusetts institute of technology"),
    ("impact-stories/index.html", "<img>", CDN + "/images/misc/aashna.png",                         "Aashna"),
    ("impact-stories/index.html", "<img>", CDN + "/images/logos/partners/cutner-elementary.png",    "massachusetts institute of technology"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/students/ammran-mohamed.png",        "Ammran H Mohamed"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/students/vitaliy-stepanov.png",      "Vitaliy Stephanov"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/students/alexa-urrea.png",           "Alexa Urrea"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/students/jonathan-williams.png",     "Jonathan Williams"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/students/netra-ramesh.png",          "Netra Ramesh"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/students/jed-rendo-m.png",           "Jed Rendo Margarcia"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/students/ashmita-kumar.png",         "Ashmita Kumar"),
    ("impact-stories/index.html", "<img>", CDN + "/images/team/students/ashna-khetan.png",          "Ashna Khetan"),
    ("impact-stories/index.html", "<img>", "https://images.squarespace-cdn.com/content/v1/5f45536caa356e6ab51588f4/1633534584521-GRMHMPU1505DA23UV5FV/Picture1.png", "Jordan Career Decision Video"),
    ("impact-stories/index.html", "<img>", CDN + "/images/misc/rezwan.jpeg",                        "Rezwan Mentorship Video"),
    ("impact-stories/index.html", "<img>", CDN + "/images/projects/makeathon-spotlight.png",        "Makeathon Spotlight Video"),
    ("impact-stories/index.html", "<img>", CDN + "/images/logos/partners/affectiva.png",            "Affectiva"),

    # ── <img> — for-students/index.html ──────────────────────────────────────
    ("for-students/index.html", "<img>", CDN + "/images/logos/surestart/surestart-logo.png",    "SureStart"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/universities/harvard.png",         "Harvard University"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/universities/uc-berkeley.png",     "UC Berkeley"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/universities/princeton.png",       "Princeton University"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/universities/stanford.png",        "Stanford University"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/universities/dartmouth.png",       "Dartmouth"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/universities/cuny.png",            "The City University of New York"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/universities/georgia-tech.png",    "Georgia Tech"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/companies/microsoft.png",          "Microsoft"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/companies/nasa.png",               "NASA"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/companies/meta.png",               "Meta"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/companies/consumer-reports.png",   "Consumer Reports"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/companies/red-hat.png",            "Red Hat"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/companies/salesforce.png",         "Salesforce"),
    ("for-students/index.html", "<img>", CDN + "/images/logos/companies/linkedin.png",           "Linkedin"),

    # ── <img> — error pages / cookies ────────────────────────────────────────
    ("cookies/index.html",      "<img>", CDN + "/images/logos/surestart/surestart-logo.png", "SureStart"),
    ("404.html",                "<img>", CDN + "/images/logos/surestart/surestart-logo.png", "SureStart"),
    ("error-pages/403.html",    "<img>", CDN + "/images/logos/surestart/surestart-logo.png", "SureStart"),
    ("error-pages/500.html",    "<img>", CDN + "/images/logos/surestart/surestart-logo.png", "SureStart"),
    ("error-pages/503.html",    "<img>", CDN + "/images/logos/surestart/surestart-logo.png", "SureStart"),
]

# ── Build workbook ────────────────────────────────────────────────────────────

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Assets"

# Colours
HEADER_FILL   = PatternFill("solid", fgColor="1F3864")   # dark navy
CSS_FILL      = PatternFill("solid", fgColor="FFF2CC")   # light yellow
INLINE_FILL   = PatternFill("solid", fgColor="FCE4D6")   # light orange
IMG_FILL      = PatternFill("solid", fgColor="E2EFDA")   # light green
VIDEO_FILL    = PatternFill("solid", fgColor="DDEBF7")   # light blue
MISSING_FILL  = PatternFill("solid", fgColor="FFD7D7")   # light red  (no alt on <img>/<video>)

thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Header
headers = ["Source File", "Type", "URL", "Alt Text"]
ws.append(headers)
for col, _ in enumerate(headers, 1):
    cell = ws.cell(1, col)
    cell.font      = Font(bold=True, color="FFFFFF", size=11)
    cell.fill      = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = border
ws.row_dimensions[1].height = 22

# Data rows
for r_idx, (source, kind, url, alt) in enumerate(rows, 2):
    ws.cell(r_idx, 1, source)
    ws.cell(r_idx, 2, kind)
    ws.cell(r_idx, 3, url)
    ws.cell(r_idx, 4, alt)

    # Row colour by type
    if kind == "CSS background":
        row_fill = CSS_FILL
    elif kind == "Inline CSS background":
        row_fill = INLINE_FILL
    elif kind == "<video>":
        row_fill = VIDEO_FILL
    else:
        row_fill = IMG_FILL

    # Flag missing alt on <img> and <video>
    if kind in ("<img>", "<video>") and alt == "":
        row_fill = MISSING_FILL

    for col in range(1, 5):
        cell = ws.cell(r_idx, col)
        cell.fill      = row_fill
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        cell.border    = border

# Column widths
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 90
ws.column_dimensions["D"].width = 48

# Freeze header row
ws.freeze_panes = "A2"

# Auto-filter
ws.auto_filter.ref = f"A1:D{len(rows)+1}"

# Legend sheet
lg = wb.create_sheet("Legend")
legend = [
    ("Colour",           "Meaning"),
    ("Dark navy header", "Column headers"),
    ("Light green",      "<img> tag with alt text"),
    ("Light red",        "<img> or <video> tag MISSING alt text"),
    ("Light blue",       "<video> tag"),
    ("Light yellow",     "CSS background-image (stylesheet)"),
    ("Light orange",     "Inline style background-image (HTML attribute)"),
]
for r, (label, meaning) in enumerate(legend, 1):
    lg.cell(r, 1, label)
    lg.cell(r, 2, meaning)
    if r == 1:
        lg.cell(r, 1).font = Font(bold=True)
        lg.cell(r, 2).font = Font(bold=True)
lg.column_dimensions["A"].width = 28
lg.column_dimensions["B"].width = 52

out = r"d:\surestart-website\assets-audit.xlsx"
wb.save(out)
print(f"Saved: {out}  ({len(rows)} rows)")
